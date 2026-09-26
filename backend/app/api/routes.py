from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from functools import wraps
import csv
from io import BytesIO, StringIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from flask import Blueprint, Response, jsonify, request, session
from sqlalchemy import func, text
from werkzeug.security import check_password_hash, generate_password_hash

from ..extensions import db
from ..models import Budget, Category, Transaction, User

api = Blueprint("api", __name__)
DEFAULT_CATEGORY_NAMES = ("Food", "Transport", "Bills", "Shopping", "Health", "Entertainment")


def current_user():
    user_id = session.get("user_id")
    return db.session.get(User, user_id) if user_id else None


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not current_user():
            return jsonify({"error": "authentication required"}), 401
        return view(*args, **kwargs)

    return wrapped


@api.get("/health")
def health_check():
    db.session.execute(text("SELECT 1"))
    return jsonify({"status": "ok", "database": "connected"})


@api.get("/v1/status")
def api_status():
    return jsonify({"name": "expanse-api", "version": "v1"})


@api.post("/v1/auth/register")
def register():
    payload = request.get_json(silent=True) or {}
    email = str(payload.get("email", "")).strip().lower()
    password = str(payload.get("password", ""))
    display_name = str(payload.get("display_name", "")).strip() or None

    if "@" not in email or len(email) > 255:
        return jsonify({"error": "enter a valid email address"}), 400
    if len(password) < 8:
        return jsonify({"error": "password must be at least 8 characters"}), 400
    if User.query.filter_by(email=email).first():
        return jsonify({"error": "an account with that email already exists"}), 409

    user = User(email=email, password_hash=generate_password_hash(password), display_name=display_name)
    db.session.add(user)
    db.session.commit()
    session["user_id"] = user.id
    return jsonify(user_payload(user)), 201


@api.post("/v1/auth/login")
def login():
    payload = request.get_json(silent=True) or {}
    email = str(payload.get("email", "")).strip().lower()
    password = str(payload.get("password", ""))
    user = User.query.filter_by(email=email).first()

    if not user or not check_password_hash(user.password_hash, password):
        return jsonify({"error": "invalid email or password"}), 401

    session.clear()
    session["user_id"] = user.id
    return jsonify(user_payload(user))


@api.patch("/v1/auth/profile")
@login_required
def update_profile():
    user = current_user()
    payload = request.get_json(silent=True) or {}
    display_name = str(payload.get("display_name", user.display_name or "")).strip()
    theme = str(payload.get("theme", user.theme)).lower()
    if not display_name or len(display_name) > 80:
        return jsonify({"error": "name must be between 1 and 80 characters"}), 400
    if theme not in {"light", "dark"}:
        return jsonify({"error": "theme must be light or dark"}), 400
    user.display_name = display_name
    user.theme = theme
    db.session.commit()
    return jsonify(user_payload(user))


@api.post("/v1/auth/logout")
def logout():
    session.clear()
    return "", 204


@api.get("/v1/auth/me")
def me():
    user = current_user()
    if not user:
        return jsonify({"error": "authentication required"}), 401
    return jsonify({"id": user.id, "email": user.email, "display_name": user.display_name, "theme": user.theme})


def user_payload(user):
    return {"id": user.id, "email": user.email, "display_name": user.display_name, "theme": user.theme}


@api.get("/v1/categories")
@login_required
def list_categories():
    existing_names = {category.name for category in Category.query.all()}
    for name in DEFAULT_CATEGORY_NAMES:
        if name not in existing_names:
            db.session.add(Category(name=name))
    db.session.commit()
    categories = Category.query.order_by(Category.name).all()
    return jsonify([{"id": category.id, "name": category.name} for category in categories])


@api.post("/v1/categories")
@login_required
def create_category():
    payload = request.get_json(silent=True) or {}
    name = str(payload.get("name", "")).strip()

    if not name:
        return jsonify({"error": "name is required"}), 400
    if Category.query.filter_by(name=name).first():
        return jsonify({"error": "category already exists"}), 409

    category = Category(name=name)
    db.session.add(category)
    db.session.commit()
    return jsonify({"id": category.id, "name": category.name}), 201


@api.get("/v1/transactions")
@login_required
def list_transactions():
    transactions = Transaction.query.filter_by(user_id=current_user().id).order_by(Transaction.created_at.desc()).all()
    return jsonify([transaction.to_dict() for transaction in transactions])


@api.delete("/v1/transactions/<int:transaction_id>")
@login_required
def delete_transaction(transaction_id):
    transaction = Transaction.query.filter_by(id=transaction_id, user_id=current_user().id).first()
    if not transaction:
        return jsonify({"error": "transaction not found"}), 404

    db.session.delete(transaction)
    db.session.commit()
    return "", 204


@api.get("/v1/dashboard/summary")
@login_required
def dashboard_summary():
    user_id = current_user().id
    income = db.session.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
        Transaction.transaction_type == "income", Transaction.user_id == user_id
    ).scalar()
    expenses = db.session.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
        Transaction.transaction_type == "expense", Transaction.user_id == user_id
    ).scalar()
    category_totals = db.session.query(
        Category.name,
        func.sum(Transaction.amount),
    ).join(Transaction, Transaction.category_id == Category.id).filter(
        Transaction.transaction_type == "expense", Transaction.user_id == user_id
    ).group_by(Category.name).order_by(func.sum(Transaction.amount).desc()).all()

    return jsonify({
        "balance": f"{income - expenses:.2f}",
        "income": f"{income:.2f}",
        "expenses": f"{expenses:.2f}",
        "transaction_count": Transaction.query.filter_by(user_id=user_id).count(),
        "category_totals": [
            {"category": name, "amount": f"{amount:.2f}"}
            for name, amount in category_totals
        ],
    })


@api.get("/v1/reports")
@login_required
def reports():
    transactions = Transaction.query.filter_by(user_id=current_user().id).order_by(Transaction.transaction_date).all()
    monthly = {}
    categories = {}
    for transaction in transactions:
        month = transaction.transaction_date.strftime("%Y-%m")
        bucket = monthly.setdefault(month, {"month": month, "income": Decimal("0"), "expenses": Decimal("0")})
        if transaction.transaction_type == "income":
            bucket["income"] += Decimal(transaction.amount)
        else:
            bucket["expenses"] += Decimal(transaction.amount)
            category = transaction.category.name if transaction.category else "Uncategorized"
            categories[category] = categories.get(category, Decimal("0")) + Decimal(transaction.amount)

    monthly_rows = []
    for bucket in monthly.values():
        bucket["net"] = bucket["income"] - bucket["expenses"]
        monthly_rows.append({key: (f"{value:.2f}" if isinstance(value, Decimal) else value) for key, value in bucket.items()})

    category_rows = [
        {"category": name, "amount": f"{amount:.2f}"}
        for name, amount in sorted(categories.items(), key=lambda item: item[1], reverse=True)
    ]
    return jsonify({"monthly": monthly_rows, "categories": category_rows})


@api.get("/v1/transactions/export")
@login_required
def export_transactions():
    transaction_query = Transaction.query.filter_by(user_id=current_user().id)
    export_type = request.args.get("type")
    if export_type in {"income", "expense"}:
        transaction_query = transaction_query.filter_by(transaction_type=export_type)
    transactions = transaction_query.order_by(Transaction.created_at.desc()).all()
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["date", "type", "description", "category", "amount"])
    for transaction in transactions:
        writer.writerow([
            transaction.transaction_date.isoformat(),
            transaction.transaction_type,
            transaction.description,
            transaction.category.name if transaction.category else "Uncategorized",
            f"{Decimal(transaction.amount):.2f}",
        ])
    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=expanse-{export_type or 'transactions'}.csv"},
    )


@api.get("/v1/transactions/export.xlsx")
@login_required
def export_transactions_xlsx():
    transaction_query = Transaction.query.filter_by(user_id=current_user().id)
    export_type = request.args.get("type")
    if export_type in {"income", "expense"}:
        transaction_query = transaction_query.filter_by(transaction_type=export_type)
    transactions = transaction_query.order_by(Transaction.created_at.desc()).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    headers = ["Date", "Type", "Description", "Category", "Amount"]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="F0B45D")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="14221F")

    for transaction in transactions:
        sheet.append([
            transaction.transaction_date.isoformat(),
            transaction.transaction_type.title(),
            transaction.description,
            transaction.category.name if transaction.category else "Uncategorized",
            float(transaction.amount),
        ])
        fill = PatternFill("solid", fgColor="E5F5EC" if transaction.transaction_type == "income" else "FBE5E2")
        for cell in sheet[sheet.max_row]:
            cell.fill = fill

    for index in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(index)].width = 50
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=expanse-{export_type or 'transactions'}.xlsx"},
    )


def parse_month(value):
    return datetime.strptime(value, "%Y-%m").date().replace(day=1)


def budget_with_spent(budget):
    spent = db.session.query(func.coalesce(func.sum(Transaction.amount), 0)).filter(
        Transaction.user_id == current_user().id,
        Transaction.category_id == budget.category_id,
        Transaction.transaction_type == "expense",
        Transaction.transaction_date >= budget.month,
        Transaction.transaction_date < date(
            budget.month.year + (budget.month.month == 12),
            1 if budget.month.month == 12 else budget.month.month + 1,
            1,
        ),
    ).scalar()
    return budget.to_dict(spent)


@api.get("/v1/budgets")
@login_required
def list_budgets():
    budgets = Budget.query.filter_by(user_id=current_user().id).order_by(Budget.month.desc()).all()
    return jsonify([budget_with_spent(budget) for budget in budgets])


@api.post("/v1/budgets")
@login_required
def create_budget():
    payload = request.get_json(silent=True) or {}
    try:
        amount = Decimal(str(payload.get("amount")))
        month = parse_month(str(payload.get("month", "")))
        category_id = int(payload.get("category_id"))
    except (InvalidOperation, TypeError, ValueError):
        return jsonify({"error": "amount, category, and month are required"}), 400
    if not amount.is_finite() or amount <= 0:
        return jsonify({"error": "budget amount must be positive"}), 400
    if not db.session.get(Category, category_id):
        return jsonify({"error": "category not found"}), 404
    if Budget.query.filter_by(user_id=current_user().id, category_id=category_id, month=month).first():
        return jsonify({"error": "a budget already exists for this category and month"}), 409

    budget = Budget(amount=amount, month=month, category_id=category_id, user_id=current_user().id)
    db.session.add(budget)
    db.session.commit()
    return jsonify(budget_with_spent(budget)), 201


@api.delete("/v1/budgets/<int:budget_id>")
@login_required
def delete_budget(budget_id):
    budget = Budget.query.filter_by(id=budget_id, user_id=current_user().id).first()
    if not budget:
        return jsonify({"error": "budget not found"}), 404
    db.session.delete(budget)
    db.session.commit()
    return "", 204


@api.post("/v1/transactions")
@login_required
def create_transaction():
    payload = request.get_json(silent=True) or {}
    transaction_type = payload.get("type")
    description = str(payload.get("description", "")).strip()

    if transaction_type not in {"income", "expense"}:
        return jsonify({"error": "type must be income or expense"}), 400
    if not description:
        return jsonify({"error": "description is required"}), 400

    try:
        amount = Decimal(str(payload.get("amount")))
    except (InvalidOperation, TypeError):
        return jsonify({"error": "amount must be a positive number"}), 400
    if not amount.is_finite() or amount <= 0:
        return jsonify({"error": "amount must be a positive number"}), 400

    try:
        transaction_date = date.fromisoformat(payload.get("date", date.today().isoformat()))
    except (TypeError, ValueError):
        return jsonify({"error": "date must use YYYY-MM-DD format"}), 400

    category_id = payload.get("category_id")
    if category_id is not None and not db.session.get(Category, category_id):
        return jsonify({"error": "category not found"}), 404

    transaction = Transaction(
        amount=amount,
        transaction_type=transaction_type,
        description=description,
        transaction_date=transaction_date,
        user_id=current_user().id,
        category_id=category_id,
    )
    db.session.add(transaction)
    db.session.commit()
    return jsonify(transaction.to_dict()), 201


@api.put("/v1/transactions/<int:transaction_id>")
@login_required
def update_transaction(transaction_id):
    transaction = Transaction.query.filter_by(id=transaction_id, user_id=current_user().id).first()
    if not transaction:
        return jsonify({"error": "transaction not found"}), 404

    payload = request.get_json(silent=True) or {}
    transaction_type = payload.get("type")
    description = str(payload.get("description", "")).strip()
    if transaction_type not in {"income", "expense"}:
        return jsonify({"error": "type must be income or expense"}), 400
    if not description:
        return jsonify({"error": "description is required"}), 400

    try:
        amount = Decimal(str(payload.get("amount")))
        transaction_date = date.fromisoformat(payload.get("date", date.today().isoformat()))
    except (InvalidOperation, TypeError, ValueError):
        return jsonify({"error": "amount or date is invalid"}), 400
    if not amount.is_finite() or amount <= 0:
        return jsonify({"error": "amount must be a positive number"}), 400

    category_id = payload.get("category_id")
    if category_id is not None and not db.session.get(Category, category_id):
        return jsonify({"error": "category not found"}), 404

    transaction.amount = amount
    transaction.transaction_type = transaction_type
    transaction.description = description
    transaction.transaction_date = transaction_date
    transaction.category_id = category_id
    db.session.commit()
    return jsonify(transaction.to_dict())
